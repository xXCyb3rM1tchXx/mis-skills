#!/usr/bin/env python3
"""Check a structural subset of Stage 11 DOCX styling; not rendered visual QA."""
import sys
from docx import Document
from docx.oxml.ns import qn

EXPECTED = {
    "Normal": ("Aptos", 9.5, None),
    "Title": ("Aptos Display", 22.0, "1F4E79"),
    "Heading 1": ("Aptos Display", 15.0, "1F4E79"),
    "Heading 2": ("Aptos Display", 12.0, "2F5597"),
    "Heading 3": ("Aptos Display", 10.5, "4472C4"),
}

def rgb(font):
    try:
        return str(font.color.rgb) if font.color.rgb else None
    except Exception:
        return None

def close(a,b,tol=.03):
    return abs(a-b)<=tol

def main(path):
    doc=Document(path)
    errors=[]
    for section_number, sec in enumerate(doc.sections, 1):
        prefix = f"Section {section_number}"
        dims = tuple(v.inches if v is not None else None
                     for v in (sec.page_width, sec.page_height))
        if any(v is None for v in dims) or not (
                close(dims[0], 8.5) and close(dims[1], 11.0)):
            errors.append(f"{prefix} page size {dims} != Letter")
        for label, value, expected in [
            ("top", sec.top_margin, .55), ("bottom", sec.bottom_margin, .55),
            ("left", sec.left_margin, .62), ("right", sec.right_margin, .62)
        ]:
            if value is None or not close(value.inches, expected, .04):
                errors.append(f"{prefix} {label} margin missing or != ~{expected}")
        variants = [("default", sec.header, sec.footer)]
        if sec.different_first_page_header_footer:
            variants.append(("first page", sec.first_page_header, sec.first_page_footer))
        if doc.settings.odd_and_even_pages_header_footer:
            variants.append(("even page", sec.even_page_header, sec.even_page_footer))
        for variant, header_part, footer_part in variants:
            header = ' '.join(p.text for p in header_part.paragraphs).strip()
            footer = ' '.join(p.text for p in footer_part.paragraphs).strip()
            if 'Prospect Intelligence Book' not in header:
                errors.append(f'{prefix} {variant} header missing Prospect Intelligence Book')
            if 'Uso interno comercial' not in footer:
                errors.append(f'{prefix} {variant} footer missing Uso interno comercial')
    for name,(font,size,color) in EXPECTED.items():
        if name not in doc.styles:
            errors.append(f"Required style {name} missing")
            continue
        st=doc.styles[name]
        f=st.font
        if f.name != font:
            errors.append(f"{name} font {f.name} != {font}")
        if f.size is None or abs(f.size.pt-size)>.1:
            errors.append(f"{name} size {f.size.pt if f.size is not None else None} != {size}")
        if color and rgb(f)!=color:
            errors.append(f"{name} color {rgb(f)} != {color}")
    # Title/headings cannot override the mandated hierarchy via direct formatting.
    # Normal also serves account names, notes and table cells with intentional
    # size/color variations; leave those roles to rendered visual inspection.
    for paragraph_number, paragraph in enumerate(doc.paragraphs, 1):
        name = paragraph.style.name
        if name not in EXPECTED or name == "Normal":
            continue
        font, size, color = EXPECTED[name]
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            formatting = [run.font]
            if run.style is not None:
                style = run.style
                while style is not None:
                    formatting.append(style.font)
                    style = style.base_style
            formatting.append(paragraph.style.font)
            # Resolve the highest-precedence explicit property. Unset run
            # properties legitimately inherit the required paragraph style.
            actual_font = next((f.name for f in formatting if f.name is not None), None)
            actual_size = next((f.size for f in formatting if f.size is not None), None)
            actual_color = next((rgb(f) for f in formatting if f.color.type is not None), None)
            if actual_font != font:
                errors.append(f"Paragraph {paragraph_number} {name} font override != {font}")
            if actual_size is None or abs(actual_size.pt - size) > .1:
                errors.append(f"Paragraph {paragraph_number} {name} size override != {size}")
            if actual_color != color:
                errors.append(f"Paragraph {paragraph_number} {name} color override != {color}")
    for i,t in enumerate(doc.tables):
        if not t.rows:
            continue
        for c in t.rows[0].cells:
            shd=c._tc.get_or_add_tcPr().find(qn('w:shd'))
            fill=shd.get(qn('w:fill')) if shd is not None else None
            if fill is None or fill.upper()!='1F4E79':
                errors.append(f"Table {i} header fill {fill} != 1F4E79")
    if errors:
        print('FAIL: VISUAL_STRUCTURE_CHECK')
        for e in errors:
            print('-',e)
        raise SystemExit(1)
    print('PASS: VISUAL_STRUCTURE_CHECK')
    print('VISUAL_SYSTEM_GATE: PENDING_RENDERED_PAGE_INSPECTION')

if __name__=='__main__':
    if len(sys.argv)!=2:
        print('usage: validate_handoff_visual.py report.docx')
        raise SystemExit(2)
    main(sys.argv[1])

